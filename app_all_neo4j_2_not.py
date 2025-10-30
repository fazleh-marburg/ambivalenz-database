from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename
import os
import spacy
import csv
import random
import subprocess  # <-- add this at the top of your file with imports


# --- Configuration ---
number = random.randint(1, 100)
category = "article"
objekt_id = f"{category}_{number}"
DIR = "/home/melahi/code/A-mediawiki-project/neo4j-upload/dataset/german/input/"
CSV_FILE = os.path.join(DIR, f"entity_{objekt_id}.csv")
EXCEL_FILE = "objekt_data.xlsx"

# Initialize Flask
app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Load spaCy German model
nlp = spacy.load("de_core_news_sm")

# Ensure Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity"])
    wb.save(EXCEL_FILE)


def get_field_list():
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        properties = [str(row[0]).strip() for row in ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]]
        if properties:
            return properties
    return [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]


def extract_name_entities(text):
    if not text:
        return []
    doc = nlp(text)
    entities = [ent.text.strip() for ent in doc.ents if ent.label_ in ("PER", "ORG", "LOC")]
    seen = []
    for e in entities:
        if e and e not in seen:
            seen.append(e)
    return seen


@app.route('/')
def form():
    fields = get_field_list()
    prefill = request.args.get('prefill', '').strip()
    data = {fields[0]: prefill} if prefill else {}
    colors = {}
    name_entities = {}
    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors, name_entities=name_entities)


@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files or request.files['file'].filename == '':
        return "❌ Keine Datei ausgewählt", 400
    file = request.files['file']
    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)
    wb = load_workbook(filepath)
    ws = wb.active
    data = {}
    colors = {}
    name_entities = {}
    properties = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        value = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        data[key] = value
        colors[key] = "yellow" if value else "red"
        name_entities[key] = extract_name_entities(value)
        properties.append(key)
    return render_template("objekt_form_buttons.html",
                           fields=properties, data=data, colors=colors, name_entities=name_entities)


@app.route('/extract_entities', methods=['POST'])
def extract_entities():
    fields = get_field_list()
    data, colors, name_entities = {}, {}, {}
    for key in fields:
        value = request.form.get(key, "").strip()
        data[key] = value
        ne = extract_name_entities(value)
        name_entities[key] = ne
        user_color = request.form.get(f"traffic_{key}")
        colors[key] = user_color if user_color else "green" if ne else "yellow" if value else "red"
    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors, name_entities=name_entities)


@app.route('/submit', methods=['POST'])
def submit():
    """Save form data to Excel and CSV, then run Java import."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity", "nodeType"])
    ws.append(["Objekt-ID", objekt_id, "", "", category])

    fields = get_field_list()
    posted = request.form
    rows = []

    # Form fields
    for field in fields:
        val = posted.get(field, "").strip() or "No data"
        status = posted.get(f"traffic_{field}", "red")
        ne = extract_name_entities(val)
        row = [field, val, status, ", ".join(ne), category]
        ws.append(row)
        rows.append(row)

    # Custom fields
    for key, value in posted.items():
        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            prop_name = value.strip()
            prop_val = posted.get(f"custom_value_{index}", "").strip() or "No data"
            prop_status = posted.get(f"traffic_custom_{index}", "red")
            ne = extract_name_entities(prop_val)
            if prop_name:
                row = [prop_name, prop_val, prop_status, ", ".join(ne), category]
                ws.append(row)
                rows.append(row)

    # Save Excel
    wb.save(EXCEL_FILE)

    # Save CSV
    os.makedirs(os.path.dirname(CSV_FILE), exist_ok=True)
    with open(CSV_FILE, mode="w", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(["Property", "text", "Status", "Name Entity", "nodeType"])
        writer.writerow(["Objekt-ID", objekt_id, "", "", category])
        writer.writerows(rows)
        writer.writerow(["nodeType", category])

    # 🔹 Run Java program to add data to Neo4j
    try:
        result = subprocess.run(
            ["java", "-jar", "target/QuestionGrammarGenerator.jar", "CREATE"],
            capture_output=True, text=True, check=True
        )
        print("✅ Java output:", result.stdout)
    except subprocess.CalledProcessError as e:
        print("❌ Java error:", e.stderr)
        return f"<h3>⚠️ Fehler beim Ausführen des Java-Befehls:<br>{e.stderr}</h3>"

    return redirect(url_for('success'))


@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich gespeichert, als CSV/Excel exportiert!</h3>"


if __name__ == '__main__':
    app.run(debug=True, port=5010)

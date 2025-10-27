from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename
import os
import csv
import random
import subprocess  # <-- Added to run Java command

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

number = random.randint(1, 100)
category = "article"
DIR = "/home/melahi/code/A-mediawiki-project/neo4j-upload/dataset/german/input/"
CSV_FILE = DIR + "entity_" + category + "_" + str(number) + ".csv"  # output file
EXCEL_FILE = "objekt_data.xlsx"  # Input Excel file

# Ensure the base Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status"])
    wb.save(EXCEL_FILE)


def get_field_list():
    return [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]


@app.route('/')
def form():
    fields = get_field_list()
    return render_template("objekt_form_all.html", fields=fields, data={}, colors={})


@app.route('/upload', methods=['POST'])
def upload():
    """Upload Excel and pre-fill form"""
    if 'file' not in request.files:
        return "❌ Keine Datei ausgewählt", 400
    file = request.files['file']
    if file.filename == '':
        return "❌ Ungültige Datei", 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    wb = load_workbook(filepath)
    ws = wb.active

    data = {}
    colors = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        value = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        status = str(row[2]).strip() if len(row) > 2 and row[2] else "red"
        data[key] = value
        colors[key] = status

    fields = get_field_list()
    return render_template("objekt_form.html", fields=fields, data=data, colors=colors)


@app.route('/submit', methods=['POST'])
def submit():
    """Save form data to CSV (with 'X' for empty fields and nodeType row), then run Java importer"""
    with open(CSV_FILE, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)

        # Write header row
        writer.writerow(["Property", "Value", "Status"])

        # Write form data
        for key, value in request.form.items():
            if key.startswith("custom_property_"):
                index = key.split("_")[-1]
                prop_name = value.strip() or "X"
                prop_val = request.form.get(f"custom_value_{index}", "").strip() or "X"
                prop_status = request.form.get(f"traffic_custom_{index}", "red") or "X"
                writer.writerow([prop_name, prop_val, prop_status])
            elif key.startswith("traffic_"):
                continue
            else:
                prop_status = request.form.get(f"traffic_{key}", "red") or "X"
                value = value.strip() or "X"
                writer.writerow([key, value, prop_status])

        # Add the final fixed row
        writer.writerow(["nodeType", category, "X"])

    # ✅ Run the Java import command AFTER saving the CSV
    try:
        result = subprocess.run(
            ["java", "-jar", "/home/melahi/code/A-mediawiki-project/neo4j-upload/target/QuestionGrammarGenerator.jar", "CREATE"],
            check=True,
            capture_output=True,
            text=True
        )
        print("✅ Java import successful:\n", result.stdout)
    except subprocess.CalledProcessError as e:
        print("❌ Java import failed:\n", e.stderr)
        return f"<h3>❌ Fehler beim Import in Neo4j:<br>{e.stderr}</h3>", 500

    return redirect(url_for('success'))


@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich in CSV gespeichert und in Neo4j importiert! (Leere Felder = 'X', letzte Zeile = nodeType/article)</h3>"


if __name__ == '__main__':
    app.run(debug=True)

from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename
import os
import spacy

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

EXCEL_FILE = "objekt_data.xlsx"

# Load spaCy German model
nlp = spacy.load("de_core_news_sm")

# Ensure main Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity"])
    wb.save(EXCEL_FILE)

def get_field_list():
    return [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]

def extract_name_entities(text):
    """Return comma-separated named entities from text"""
    if not text:
        return ""
    doc = nlp(text)
    entities = [ent.text for ent in doc.ents if ent.label_ in ["PER", "ORG", "LOC"]]
    return ", ".join(entities)

@app.route('/')
def form():
    fields = get_field_list()
    return render_template("objekt_form_all.html", fields=fields, data={}, colors={}, name_entities={})

@app.route('/upload', methods=['POST'])
def upload():
    """Upload Excel and pre-fill form"""
    if 'file' not in request.files:
        return "❌ Keine Datei ausgewählt", 400
    file = request.files['file']
    if file.filename == '':
        return "❌ Ungültige Datei", 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    wb = load_workbook(filepath)
    ws = wb.active

    data = {}
    colors = {}
    name_entities = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        value = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        data[key] = value
        colors[key] = "yellow" if value else "red"
        name_entities[key] = extract_name_entities(value)

    fields = get_field_list()
    return render_template("objekt_form_all.html", fields=fields, data=data, colors=colors, name_entities=name_entities)

@app.route('/extract_entities', methods=['POST'])
def extract_entities():
    """Extract named entities from all text fields"""
    fields = get_field_list()
    data = {}
    colors = {}
    name_entities = {}

    for key in fields:
        value = request.form.get(key, "").strip()
        data[key] = value
        ne = extract_name_entities(value)
        name_entities[key] = ne
        colors[key] = "green" if ne else "yellow" if value else "red"

    return render_template("objekt_form_all.html", fields=fields, data=data, colors=colors, name_entities=name_entities)

@app.route('/submit', methods=['POST'])
def submit():
    """Save form data to Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity"])

    for key, value in request.form.items():
        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            prop_name = value.strip()
            prop_val = request.form.get(f"custom_value_{index}", "").strip()
            prop_status = request.form.get(f"traffic_custom_{index}", "red")
            if prop_name:
                ne = extract_name_entities(prop_val)
                ws.append([prop_name, prop_val, prop_status, ne])
        elif key.startswith("traffic_"):
            continue
        else:
            prop_status = request.form.get(f"traffic_{key}", "red")
            ne = extract_name_entities(value)
            ws.append([key, value, prop_status, ne])

    wb.save(EXCEL_FILE)
    return redirect(url_for('success'))

@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich gespeichert!</h3>"

if __name__ == '__main__':
    app.run(debug=True)

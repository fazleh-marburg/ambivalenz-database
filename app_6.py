from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename
import os

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

EXCEL_FILE = "objekt_data.xlsx"

# Ensure Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status"])
    wb.save(EXCEL_FILE)

def get_field_list():
    return [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]

@app.route('/')
def form():
    fields = get_field_list()
    return render_template("objekt_form.html", fields=fields, data={}, colors={})

@app.route('/upload', methods=['POST'])
def upload():
    """Upload Excel and pre-fill form"""
    if 'file' not in request.files:
        return "Keine Datei ausgewählt", 400
    file = request.files['file']
    if file.filename == '':
        return "Ungültige Datei", 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    wb = load_workbook(filepath)
    ws = wb.active

    data = {}
    colors = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        value = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        status = str(row[2]).strip() if len(row) > 2 and row[2] else "yellow"
        data[key] = value
        colors[key] = status

    fields = get_field_list()
    return render_template("objekt_form.html", fields=fields, data=data, colors=colors)

@app.route('/submit', methods=['POST'])
def submit():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    # Clear previous entries (except header)
    ws.delete_rows(2, ws.max_row)

    for key, value in request.form.items():
        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            property_name = value
            property_value = request.form.get(f"custom_value_{index}", "")
            property_status = request.form.get(f"traffic_custom_{index}", "red")
            if property_name.strip():
                ws.append([property_name, property_value, property_status])
        elif key.startswith("traffic_"):
            continue  # skip hidden traffic inputs
        else:
            property_status = request.form.get(f"traffic_{key}", "red")
            ws.append([key, value, property_status])

    wb.save(EXCEL_FILE)
    return redirect(url_for('success'))

@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich gespeichert!</h3>"

if __name__ == '__main__':
    app.run(debug=True)

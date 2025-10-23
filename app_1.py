from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
import os

app = Flask(__name__)

EXCEL_FILE = "objekt_data.xlsx"

# Ensure Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value"])
    wb.save(EXCEL_FILE)

@app.route('/')
def form():
    # Define fields (German labels as in your example)
    fields = [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]
    return render_template("objekt_form.html", fields=fields)

@app.route('/submit', methods=['POST'])
def submit():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    for key, value in request.form.items():
        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            property_name = value
            property_value = request.form.get(f"custom_value_{index}", "")
            if property_name.strip():
                ws.append([property_name, property_value])
        else:
            ws.append([key, value])

    wb.save(EXCEL_FILE)
    return redirect(url_for('success'))

@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich gespeichert!</h3>"

if __name__ == '__main__':
    app.run(debug=True)

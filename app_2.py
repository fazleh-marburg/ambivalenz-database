from flask import Flask, render_template, request, redirect, url_for, send_file
from openpyxl import Workbook, load_workbook
import os
from werkzeug.utils import secure_filename

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

EXCEL_FILE = "objekt_data.xlsx"

# Ensure Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value"])
    wb.save(EXCEL_FILE)

@app.route('/')
def form():
    # German fields
    fields = [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]
    return render_template("objekt_form.html", fields=fields, data={})

@app.route('/upload', methods=['POST'])
def upload():
    """Upload and read Excel file to prefill the form"""
    if 'file' not in request.files:
        return "❌ Keine Datei ausgewählt", 400

    file = request.files['file']
    if file.filename == '':
        return "❌ Ungültige Datei", 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    # Read Excel content
    wb = load_workbook(filepath)
    ws = wb.active

    data = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1]:
            data[row[0]] = row[1]

    # Same field list
    fields = [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]

    return render_template("objekt_form_3.html", fields=fields, data=data)

@app.route('/submit', methods=['POST'])
def submit():
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active

    for key, value in request.form.items():
        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            property_name = value
            property_value = request.form.get(f"custom_value_{index}", "")
            if property_name.strip():
                ws.append([property_name, property_value])
        else:
            ws.append([key, value])

    wb.save(EXCEL_FILE)
    return redirect(url_for('success'))

@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich gespeichert!</h3>"

if __name__ == '__main__':
    app.run(debug=True)

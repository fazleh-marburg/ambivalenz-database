from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename
from neo4j import GraphDatabase
import os
import spacy
import csv
import random

# --- Configuration ---
number = random.randint(1, 100)
category = "article"
DIR = "/home/melahi/code/A-mediawiki-project/neo4j-upload/dataset/german/input/"
objekt_id=category+"_"+str(number)
CSV_FILE = os.path.join(DIR, f"entity_{objekt_id}.csv")
EXCEL_FILE = "objekt_data.xlsx"

# Neo4j connection info
NEO4J_URI = "bolt://localhost:7687"
NEO4J_USER = "neo4j"
NEO4J_PASSWORD = "password"

# Initialize Flask
app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Connect to Neo4j
driver = GraphDatabase.driver(NEO4J_URI, auth=(NEO4J_USER, NEO4J_PASSWORD))

# Load spaCy German model
nlp = spacy.load("de_core_news_sm")

# Ensure Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity"])
    wb.save(EXCEL_FILE)


def get_field_list():
    """Read 'Property' column from Excel file, or use default list."""
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        properties = []
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            val = str(row[0]).strip() if row[0] else ""
            if val:
                properties.append(val)
        if properties:
            return properties
    return [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]


def extract_name_entities(text):
    """Extract PER, ORG, LOC entities."""
    if not text:
        return []
    doc = nlp(text)
    entities = [ent.text.strip() for ent in doc.ents if ent.label_ in ("PER", "ORG", "LOC")]
    seen = []
    for e in entities:
        if e and e not in seen:
            seen.append(e)
    return seen


def add_data_to_neo4j(rows):
    """Create ONE node with all Property–Value pairs as attributes."""
    properties = {}
    for prop, val, status, ne in rows:
        if prop and val:
            # Sanitize property name (remove spaces, slashes, etc.)
            safe_prop = prop.replace(" ", "_").replace("/", "_").replace("*", "")
            properties[safe_prop] = val

    with driver.session() as session:
        query = f"""
        CREATE (o:Objekt)
        SET o += $props
        RETURN id(o) AS node_id
        """
        result = session.run(query, props=properties)
        node_id = result.single()["node_id"]
        print(f"✅ Created Neo4j node with ID: {node_id}")
    return node_id


@app.route('/')
def form():
    fields = get_field_list()
    prefill = request.args.get('prefill', '').strip()
    data = {}
    if prefill:
        data[fields[0]] = prefill
    colors = {}
    name_entities = {}
    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors, name_entities=name_entities)


@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files or request.files['file'].filename == '':
        return "❌ Keine Datei ausgewählt", 400

    file = request.files['file']
    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    wb = load_workbook(filepath)
    ws = wb.active

    data = {}
    colors = {}
    name_entities = {}
    properties = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        value = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        data[key] = value
        colors[key] = "yellow" if value else "red"
        name_entities[key] = extract_name_entities(value)
        properties.append(key)

    fields = properties

    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors, name_entities=name_entities)


@app.route('/extract_entities', methods=['POST'])
def extract_entities():
    fields = get_field_list()
    data = {}
    colors = {}
    name_entities = {}

    for key in fields:
        value = request.form.get(key, "").strip()
        data[key] = value
        ne = extract_name_entities(value)
        name_entities[key] = ne
        user_color = request.form.get(f"traffic_{key}")
        if user_color:
            colors[key] = user_color
        else:
            colors[key] = "green" if ne else "yellow" if value else "red"

    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors, name_entities=name_entities)


@app.route('/submit', methods=['POST'])
def submit():
    """Save form data to Excel, CSV, and Neo4j."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity"])

    fields = get_field_list()
    posted = request.form

    rows = []

    for field in fields:
        val = posted.get(field, "").strip()
        if not val:
            val = "No data"  # 🔹 Replace empty values
        status = posted.get(f"traffic_{field}", "red")
        ne = extract_name_entities(val)
        row = [field, val, status, ", ".join(ne)]
        ws.append(row)
        rows.append(row)

    for key, value in posted.items():
        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            prop_name = value.strip()
            prop_val = posted.get(f"custom_value_{index}", "").strip()
            if not prop_val:
                prop_val = "No data"  # 🔹 Replace empty values in custom fields
            prop_status = posted.get(f"traffic_custom_{index}", "red")
            ne = extract_name_entities(prop_val)
            if prop_name:
                row = [prop_name, prop_val, prop_status, ", ".join(ne)]
                ws.append(row)
                rows.append(row)

    # Save Excel
    wb.save(EXCEL_FILE)

    # Save CSV
    os.makedirs(os.path.dirname(CSV_FILE), exist_ok=True)
    with open(CSV_FILE, mode="w", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        # Header
        writer.writerow(["Property", "text", "Status", "Name Entity", "nodeType"])
        # Objekt-ID row
        writer.writerow(["Objekt-ID", objekt_id, "", "", category])
        # All form rows
        writer.writerows(rows)
        # Final row with just nodeType and category
        writer.writerow(["nodeType", category])

    # Create one node in Neo4j with all properties
    add_data_to_neo4j(rows)

    return redirect(url_for('success'))



@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich gespeichert, als CSV/Excel exportiert und in Neo4j importiert (ein Knoten mit allen Eigenschaften)!</h3>"


if __name__ == '__main__':
    app.run(debug=True, port=5010)

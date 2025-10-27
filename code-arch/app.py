from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
import os

app = Flask(__name__)

EXCEL_FILE = "../objekt_data.xlsx"

# Ensure Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Color"])
    wb.save(EXCEL_FILE)


@app.route('/')
def form():
    # Define standard fields
    fields = [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]

    data = {}
    colors = {}

    # Load data from Excel
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                prop, val, color = row[0], row[1] or "", row[2] or "red"
                data[prop] = val
                colors[prop] = color

    return render_template("objekt_form.html", fields=fields, data=data, colors=colors)


@app.route('/submit', methods=['POST'])
def submit():
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Color"])

    for key, value in request.form.items():
        # skip color inputs — they are handled separately
        if key.startswith("color_"):
            continue

        # determine color
        color = request.form.get(f"color_{key}", "red")

        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            property_name = value
            property_value = request.form.get(f"custom_value_{index}", "")
            color_val = request.form.get(f"color_custom_value_{index}", "red")
            if property_name.strip():
                ws.append([property_name, property_value, color_val])
        else:
            ws.append([key, value, color])

    wb.save(EXCEL_FILE)
    return redirect(url_for('success'))


@app.route('/success')
def success():
    return "<h3>✅ Daten und Farben erfolgreich gespeichert!</h3>"


if __name__ == '__main__':
    app.run(debug=True)

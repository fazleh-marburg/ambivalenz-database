from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename
import os

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

EXCEL_FILE = "../objekt_data.xlsx"

# Ensure main Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value"])
    wb.save(EXCEL_FILE)


@app.route('/')
def form():
    fields = get_field_list()
    return render_template("object_form_import.html", fields=fields, data={})


@app.route('/upload', methods=['POST'])
def upload():
    """Allow user to import .xlsx file and auto-fill the form"""
    if 'file' not in request.files:
        return "❌ Keine Datei ausgewählt", 400

    file = request.files['file']
    if file.filename == '':
        return "❌ Ungültige Datei", 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    # Read Excel and load data
    data = {}
    wb = load_workbook(filepath)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        value = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        data[key] = value

    fields = get_field_list()
    return render_template("objekt_form.html", fields=fields, data=data)


@app.route('/submit', methods=['POST'])
def submit():
    """Save form data to Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value"])

    for key, value in request.form.items():
        if key.startswith("custom_property_"):
            index = key.split("_")[-1]
            prop_name = value.strip()
            prop_val = request.form.get(f"custom_value_{index}", "").strip()
            if prop_name:
                ws.append([prop_name, prop_val])
        else:
            ws.append([key, value])

    wb.save(EXCEL_FILE)
    return redirect(url_for('success'))


@app.route('/success')
def success():
    return "<h3>✅ Daten erfolgreich gespeichert!</h3>"


def get_field_list():
    return [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]


if __name__ == '__main__':
    app.run(debug=True)

from flask import Flask, render_template, request, redirect, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename
from neo4j import GraphDatabase
import os
import spacy
import csv
import random
import subprocess

# --- Configuration ---
category = "article"
DIR = "/home/melahi/code/A-mediawiki-project/neo4j-upload/dataset/german/input/"
CSV_FILE = os.path.join(DIR, f"entity_temp.csv")
EXCEL_FILE = "objekt_data.xlsx"

# Neo4j connection info
NEO4J_URI = "bolt://localhost:7687"
NEO4J_USER = "neo4j"
NEO4J_PASSWORD = "password"

# Initialize Flask
app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Connect to Neo4j
driver = GraphDatabase.driver(NEO4J_URI, auth=(NEO4J_USER, NEO4J_PASSWORD))

# Load spaCy German model
nlp = spacy.load("de_core_news_sm")

# Ensure Excel file exists
if not os.path.exists(EXCEL_FILE):
    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity"])
    wb.save(EXCEL_FILE)


def get_field_list():
    """Read 'Property' column from Excel file, or use default list."""
    if os.path.exists(EXCEL_FILE):
        wb = load_workbook(EXCEL_FILE)
        ws = wb.active
        properties = []
        for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
            val = str(row[0]).strip() if row[0] else ""
            if val:
                properties.append(val)
        if properties:
            return properties
    return [
        "Titel", "Künstler*in/Autor*in", "Sichtbare oder genannte Personen",
        "Entstehungsjahr", "Ort der Entstehung / Nutzung", "Gattung / Genre", "Technik",
        "Dimensionen", "Kurzbeschreibung", "Motive / Topoi", "Narrative / Diskurse",
        "Historischer Kontext", "Antiziganistische / Stigmatisierende Elemente",
        "Agency", "Verknüpfung", "Narrativwandel bei Medienwechsel", "Rezeptionsweg",
        "Sammlung / Archiv", "Provenienz", "Literatur",
        "Ausstellungen / Aufführungen / Veröffentlichungen", "Objekt- oder Werkteil",
        "Rechte / Lizenzen", "Digitalisat-Link/Pfad", "Metadaten-Status",
        "Erfasst von", "Erfassungsdatum", "Kommentar / Anmerkung", "Versionsgeschichte"
    ]


def extract_name_entities(text):
    """Extract PER, ORG, LOC entities."""
    if not text:
        return []
    doc = nlp(text)
    entities = [ent.text.strip() for ent in doc.ents if ent.label_ in ("PER", "ORG", "LOC")]
    seen = []
    for e in entities:
        if e and e not in seen:
            seen.append(e)
    return seen


@app.route('/')
def form():
    fields = get_field_list()
    data = {}
    colors = {}
    name_entities = {}
    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors,
                           name_entities=name_entities, show_new_button=True)


@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files or request.files['file'].filename == '':
        return "❌ Keine Datei ausgewählt", 400

    file = request.files['file']
    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    file.save(filepath)

    wb = load_workbook(filepath)
    ws = wb.active

    data = {}
    colors = {}
    name_entities = {}
    properties = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        value = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        data[key] = value
        colors[key] = "yellow" if value else "red"
        name_entities[key] = extract_name_entities(value)
        properties.append(key)

    fields = properties

    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors,
                           name_entities=name_entities, show_new_button=True)


@app.route('/extract_entities', methods=['POST'])
def extract_entities():
    fields = get_field_list()
    data = {}
    colors = {}
    name_entities = {}

    for key in fields:
        value = request.form.get(key, "").strip()
        data[key] = value
        ne = extract_name_entities(value)
        name_entities[key] = ne
        user_color = request.form.get(f"traffic_{key}")
        if user_color:
            colors[key] = user_color
        else:
            colors[key] = "green" if ne else "yellow" if value else "red"

    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors,
                           name_entities=name_entities, show_new_button=True)


@app.route('/submit', methods=['POST'])
def submit():
    """Save form data to Excel and CSV, then run Java import or relation creation."""
    source_node_id = request.form.get("source_node_id")
    relation_name = request.form.get("relation_name")

    wb = Workbook()
    ws = wb.active
    ws.title = "Objekt Informationen"
    ws.append(["Property", "Value", "Status", "Name Entity", "nodeType"])

    fields = get_field_list()
    posted = request.form
    rows = []

    for field in fields:
        val = posted.get(field, "").strip() or "No data"
        status = posted.get(f"traffic_{field}", "red")
        ne = extract_name_entities(val)
        row = [field, val, status, ", ".join(ne), category]
        ws.append(row)
        rows.append(row)

    # Save Excel and CSV
    wb.save(EXCEL_FILE)
    os.makedirs(os.path.dirname(CSV_FILE), exist_ok=True)
    with open(CSV_FILE, mode="w", newline="", encoding="utf-8") as csvfile:
        writer = csv.writer(csvfile)
        writer.writerow(["Property", "text", "Status", "Name Entity", "nodeType"])
        writer.writerows(rows)
        writer.writerow(["nodeType", category])

    # Run Java tool (CREATE or RELATION)
    try:
        if source_node_id and relation_name:
            print(f"🔗 Creating relation '{relation_name}' from node {source_node_id}")
            result = subprocess.run(
                ["java", "-jar", "/home/melahi/code/A-mediawiki-project/neo4j-upload/target/QuestionGrammarGenerator.jar", "RELATION"],
                capture_output=True, text=True, check=True
            )
        else:
            result = subprocess.run(
                ["java", "-jar", "/home/melahi/code/A-mediawiki-project/neo4j-upload/target/QuestionGrammarGenerator.jar", "CREATE"],
                capture_output=True, text=True, check=True
            )
        print("✅ Java output:", result.stdout)
    except subprocess.CalledProcessError as e:
        print("❌ Java error:", e.stderr)
        return f"<h3>⚠️ Fehler beim Ausführen des Java-Befehls:<br>{e.stderr}</h3>"

    # ✅ Stay on the same page, show success message
    fields = get_field_list()
    data = {}
    colors = {}
    name_entities = {}
    message = "✅ Daten erfolgreich gespeichert und importiert!"
    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors,
                           name_entities=name_entities,
                           message=message, show_new_button=True)


@app.route('/new_form')
def new_form():
    """Open a blank form manually when clicking 'Neues Objekt erfassen'."""
    fields = get_field_list()
    data = {}
    colors = {}
    name_entities = {}
    return render_template("objekt_form_buttons.html",
                           fields=fields, data=data, colors=colors,
                           name_entities=name_entities, show_new_button=True)


if __name__ == '__main__':
    app.run(debug=True, port=5010)
